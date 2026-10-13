"""Operations dashboard over the saved DATA pivot (no OLAP refresh required)."""
from datetime import date, datetime, time
from pathlib import Path
import re
import unicodedata

import openpyxl
import pandas as pd
import plotly.express as px
import streamlit as st

METRICS = ["Počet jobline", "Počet kusů"]


def text(value):
    normalized = unicodedata.normalize("NFKD", str(value or ""))
    return " ".join("".join(c for c in normalized if not unicodedata.combining(c)).lower().split())


def metric(value):
    label = text(value)
    if label == "pocet jobline":
        return METRICS[0]
    if label == "pocet kusu":
        return METRICS[1]
    return None


def parse_date(value):
    if isinstance(value, (datetime, date)):
        return pd.Timestamp(value).normalize()
    if isinstance(value, str):
        for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
            try:
                return pd.Timestamp(datetime.strptime(value.strip(), fmt))
            except ValueError:
                pass
    return None


def parse_hour(value):
    if isinstance(value, (datetime, time)):
        return value.hour
    try:
        number = float(str(value).strip().split(":")[0])
        if number.is_integer() and 0 <= number <= 23:
            return int(number)
    except (TypeError, ValueError):
        pass
    return None


def number(value):
    if value is None or value == "":
        return 0.0
    if isinstance(value, str):
        value = value.replace("\u00a0", "").replace(" ", "").replace(",", ".")
    result = float(value)
    if not pd.notna(result) or result < 0:
        raise ValueError(f"Neplatný počet: {value!r}")
    return result


def normalize_pivot(rows):
    """Find semantic headers, then decode each leaf column into date/hour/metric."""
    def find(predicate):
        matches = [(r, c) for r, row in enumerate(rows) for c, value in enumerate(row)
                   if predicate(text(value))]
        if len(matches) != 1:
            raise ValueError(f"Hlavičku nemožno jednoznačne nájsť ({len(matches)} výskytov).")
        return matches[0]

    creation_row, creation_col = find(lambda s: s.startswith("datum") and "vytvor" in s and "jobline" in s)
    hour_row, hour_col = find(lambda s: s.startswith("hodina") and "vytvor" in s and "jobline" in s)
    pick_label_row, _ = find(lambda s: s.startswith("datum") and "fronty" in s and "jobline" in s)
    pick_hour_label_row, _ = find(lambda s: s.startswith("hodina") and "fronty" in s and "jobline" in s)
    candidates = [(r, sum(metric(v) is not None for v in row)) for r, row in enumerate(rows)]
    metric_row, count = max(candidates, key=lambda item: item[1])
    if count < 2 or creation_row != hour_row or metric_row < creation_row:
        raise ValueError("Nenašla sa podporovaná hlavička kontingenčnej tabuľky.")
    leaf_cols = [c for c, value in enumerate(rows[metric_row]) if metric(value)]
    date_candidates = [(r, sum(parse_date(rows[r][c]) is not None for c in leaf_cols))
                       for r in range(pick_label_row + 1, metric_row)]
    if not date_candidates or max(n for _, n in date_candidates) == 0:
        raise ValueError("Nenašiel sa riadok dátumov pickovania.")
    date_row = max(date_candidates, key=lambda item: item[1])[0]
    hour_candidates = [(r, sum(parse_hour(rows[r][c]) is not None for c in leaf_cols))
                       for r in range(max(date_row, pick_hour_label_row) + 1, metric_row)]
    if not hour_candidates or max(n for _, n in hour_candidates) == 0:
        raise ValueError("Nenašiel sa riadok hodín pickovania.")
    pick_hour_row = max(hour_candidates, key=lambda item: item[1])[0]
    columns, unmapped = [], []
    current_date = current_hour = None
    for c in leaf_cols:
        d, h = parse_date(rows[date_row][c]), parse_hour(rows[pick_hour_row][c])
        if d is not None:
            current_date = d
        if h is not None:
            current_hour = h
        if current_date is None or current_hour is None:
            unmapped.append(c + 1)
        else:
            columns.append((c, current_date, current_hour, metric(rows[metric_row][c])))
    if not columns:
        raise ValueError("Žiadny stĺpec nemá dátum, hodinu a metriku.")
    totals = {m: [c for c, v in enumerate(rows[date_row])
                  if text(v) in (f"celkovo {text(m)}", f"celkem {text(m)}", f"total {text(m)}")]
              for m in METRICS}
    events, skipped, bad, creation_date = [], [], [], None
    loaded = 0
    for r in range(metric_row + 1, len(rows)):
        row = rows[r]
        if any(re.search(r"celkov|celkem|total|sucet|soucet", text(row[c])) for c in (creation_col, hour_col)):
            creation_date = None
            continue
        d = parse_date(row[creation_col])
        if d is not None:
            creation_date = d
        h = parse_hour(row[hour_col])
        if creation_date is None or h is None:
            if any(v is not None for v in row):
                skipped.append(r + 1)
            continue
        loaded += 1
        per_metric = {m: 0.0 for m in METRICS}
        for c, pick_date, pick_hour, m in columns:
            try:
                value = number(row[c])
            except (ValueError, TypeError):
                bad.append(f"{openpyxl.utils.get_column_letter(c + 1)}{r + 1}")
                continue
            per_metric[m] += value
            if value:
                events.append({"created_at": creation_date + pd.Timedelta(hours=h),
                               "picked_at": pd.NaT if pick_date <= pd.Timestamp("1900-01-01") else pick_date + pd.Timedelta(hours=pick_hour),
                               "metric": m, "value": value})
        for m, cols in totals.items():
            if len(cols) == 1:
                try:
                    if abs(number(row[cols[0]]) - per_metric[m]) > 0.001:
                        bad.append(f"Nesúlad súčtu {m}, riadok {r + 1}")
                except (ValueError, TypeError):
                    bad.append(f"Neplatný súčet {m}, riadok {r + 1}")
    detail = pd.DataFrame(events, columns=["created_at", "picked_at", "metric", "value"])
    for col in ("created_at", "picked_at"):
        detail[col] = pd.to_datetime(detail[col])
    def grouped(frame, axis):
        if frame.empty:
            return pd.DataFrame({"Čas": pd.Series(dtype="datetime64[ns]"),
                                 **{m: pd.Series(dtype="float64") for m in METRICS}})
        result = frame.groupby([axis, "metric"])["value"].sum().unstack(fill_value=0)
        return result.reindex(columns=METRICS, fill_value=0).rename_axis("Čas").reset_index()
    diagnostics = {"Načítané riadky": loaded, "Rozmer hárka": [len(rows), len(rows[0])],
                   "Riadky hlavičiek (Excel)": {"vznik": creation_row + 1, "pick dátum": date_row + 1,
                                               "pick hodina": pick_hour_row + 1, "metriky": metric_row + 1},
                   "Stĺpce vzniku (Excel)": [creation_col + 1, hour_col + 1],
                   "Namapované stĺpce pickovania": len(columns), "Nenamapované stĺpce": unmapped,
                   "Preskočené riadky": skipped, "Neplatné hodnoty / nesúlad súčtov": bad,
                   "Mapovanie": [{"stĺpec": openpyxl.utils.get_column_letter(c + 1), "dátum": str(d.date()),
                                  "hodina": h, "metrika": m} for c, d, h, m in columns]}
    return grouped(detail, "created_at"), grouped(detail.dropna(subset=["picked_at"]), "picked_at"), grouped(detail[detail.picked_at.isna()], "created_at"), diagnostics


@st.cache_data(show_spinner="Načítavam Excel…")
def load_data(path, modified_ns, size):
    # The fingerprint invalidates the cache when the source workbook changes.
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if "DATA" not in workbook.sheetnames:
            raise ValueError("V súbore chýba hárok DATA.")
        return normalize_pivot(list(workbook["DATA"].iter_rows(values_only=True)))
    finally:
        workbook.close()


def aggregate(created, picked, start, end, frequency):
    index = pd.date_range(pd.Timestamp(start), pd.Timestamp(end) + pd.Timedelta(days=1),
                          freq=frequency, inclusive="left")
    output = pd.DataFrame(index=index)
    for name, frame in (("Vzniky", created), ("Picky", picked)):
        selected = frame[(frame["Čas"] >= pd.Timestamp(start)) & (frame["Čas"] < pd.Timestamp(end) + pd.Timedelta(days=1))].copy()
        selected["Čas"] = selected["Čas"].dt.floor(frequency)
        values = selected.groupby("Čas")[METRICS].sum()
        for m, suffix in zip(METRICS, ("jobline", "kusy")):
            output[f"{name} {suffix}"] = values[m].reindex(index, fill_value=0)
    return output.rename_axis("Čas").reset_index()


def main():
    st.set_page_config(page_title="Dashboard provoz", page_icon="📊", layout="wide")
    st.title("Dashboard provoz")
    st.caption("Vznik jobline a zaradenie do fronty pickovania podľa uložených dát v Exceli.")
    source = Path(__file__).resolve().with_name("datax.xlsx")
    try:
        stat = source.stat()
        created, picked, unknown, diagnostics = load_data(str(source), stat.st_mtime_ns, stat.st_size)
    except Exception as exc:
        st.error(f"Excel sa nepodarilo načítať: {exc}")
        with st.expander("Diagnostika dát"):
            st.write({"Súbor": str(source), "Chyba": str(exc)})
        return
    with st.expander("Diagnostika dát"):
        st.json({k: v for k, v in diagnostics.items() if k != "Mapovanie"})
        st.dataframe(pd.DataFrame(diagnostics["Mapovanie"]), hide_index=True)
        for label, frame in (("Normalizované vzniky", created), ("Normalizované picky", picked), ("Bez pick dátumu (podľa vzniku)", unknown)):
            st.write(label)
            st.dataframe(frame.head(20), hide_index=True)
    if diagnostics["Nenamapované stĺpce"] or diagnostics["Neplatné hodnoty / nesúlad súčtov"] or diagnostics["Preskočené riadky"]:
        st.error("Niektoré dáta nemožno spoľahlivo namapovať. Pozrite diagnostiku; súčty a grafy sa nezobrazia, aby neboli zavádzajúce.")
        return
    if created.empty:
        st.info("V hárku DATA nie sú žiadne nenulové záznamy.")
        return
    times = pd.concat([created["Čas"], picked["Čas"]])
    first, last = times.min().date(), times.max().date()
    selection = st.sidebar.date_input("Rozsah dátumov", value=(first, last), min_value=first, max_value=last)
    aggregation = st.sidebar.selectbox("Agregácia", ["hodina", "deň"])
    st.sidebar.caption("Vzniky filtrujeme podľa času vzniku, picky podľa času fronty. Nezaradené podľa času vzniku.")
    if len(selection) != 2:
        st.info("Vyberte začiatok aj koniec rozsahu.")
        return
    start, end = selection
    combined = aggregate(created, picked, start, end, "h" if aggregation == "hodina" else "D")
    cards = st.columns(4)
    for card, label in zip(cards, combined.columns[1:]):
        card.metric(f"{label} spolu", f"{combined[label].sum():,.0f}".replace(",", " "))
    selected_unknown = unknown[(unknown["Čas"] >= pd.Timestamp(start)) & (unknown["Čas"] < pd.Timestamp(end) + pd.Timedelta(days=1))]
    for card, m in zip(st.columns(2), METRICS):
        card.metric(f"Nezaradené / bez pick dátumu – {m}", f"{selected_unknown[m].sum():,.0f}".replace(",", " "))
    st.caption("Picky označujú zaradenie do fronty pickovania, nie potvrdenie dokončeného picku. Dátum 01.01.1900 je v samostatnom KPI.")
    for suffix in ("jobline", "kusy"):
        fig = px.line(combined, x="Čas", y=[f"Vzniky {suffix}", f"Picky {suffix}"],
                      title=f"Vzniky a picky – {suffix}", labels={"value": "Počet", "variable": "Metrika"},
                      color_discrete_sequence=["#2563eb", "#f97316"])
        fig.update_layout(hovermode="x unified", legend_title_text="")
        st.plotly_chart(fig, width="stretch")
    st.subheader("Agregované dáta")
    st.dataframe(combined, hide_index=True, width="stretch")


if __name__ == "__main__":
    main()
